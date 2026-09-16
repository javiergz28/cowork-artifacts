#!/usr/bin/env python3
"""Convierte una exportación local de STOCK_MT_FINAL en un snapshot privado.

Se usa para revisar una primera corrida descargada manualmente. La automatización
de GitHub obtiene los mismos rangos directamente desde Google Sheets.
"""
from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


def values(sheet):
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Uso: xlsx-stock-snapshot.py <STOCK_MT_FINAL.xlsx> <snapshot.json>")
    source = Path(sys.argv[1]).resolve()
    destination = Path(sys.argv[2]).resolve()
    workbook = load_workbook(source, read_only=True, data_only=True)
    for tab in ("Articulos", "Parametros"):
        if tab not in workbook.sheetnames:
            raise SystemExit(f"Falta la pestaña requerida: {tab}")
    snapshot = {
        "meta": {
            "title": "STOCK_MT_FINAL",
            "spreadsheetId": "1sDn8kLM_ewCPgsYS01n9EQTUpfmNYvZreU1leYxDB6Q",
            "fetchedAtUtc": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
            "timeZone": "America/Montevideo",
            "source": "Exportación manual de Google Sheets; Articulos + Parametros"
        },
        "articles": values(workbook["Articulos"]),
        "parameters": values(workbook["Parametros"])
    }
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(snapshot, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Snapshot privado creado: {destination}")
    print(f"Filas leídas: Articulos={len(snapshot['articles'])}, Parametros={len(snapshot['parameters'])}")


if __name__ == "__main__":
    main()
